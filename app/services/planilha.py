# -*- coding: utf-8 -*-
"""
Módulo de exportação de transcrições e métricas de qualidade para planilha Excel (.xlsx).

Armazena o histórico completo de transcrições com todas as métricas
textuais, métricas do Whisper, pontuação geral e classificação de qualidade
(Excelente, Bom, Médio, Ruim, Péssimo).
"""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Diretório raiz do projeto (dois níveis acima de app/services)
_RAIZ = Path(__file__).resolve().parents[2]
CAMINHO_PLANILHA_PADRAO = _RAIZ / "transcricoes_avaliadas.xlsx"

# Cores para classificação de qualidade (fundo e texto)
CORES_CLASSIFICACAO: Dict[str, Dict[str, Any]] = {
    "Excelente": {
        "fill": PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="274E13"),
    },
    "Bom": {
        "fill": PatternFill(start_color="D0E0E3", end_color="D0E0E3", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="0C343D"),
    },
    "Médio": {
        "fill": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="7F6000"),
    },
    "Ruim": {
        "fill": PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="783F04"),
    },
    "Péssimo": {
        "fill": PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid"),
        "font": Font(name="Calibri", size=11, bold=True, color="990000"),
    },
}

# Estilos de cabeçalho
ESTILO_CABECALHO_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ESTILO_CABECALHO_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BORDA_FINA = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)

COLUNAS = [
    ("Data / Hora", 18, Alignment(horizontal="center", vertical="top")),
    ("Arquivo de Áudio", 35, Alignment(horizontal="left", vertical="top")),
    ("Qualidade Medida", 16, Alignment(horizontal="center", vertical="top")),
    ("Pontuação", 12, Alignment(horizontal="center", vertical="top")),
    ("Status no Fluxo", 22, Alignment(horizontal="center", vertical="top")),
    ("Total Palavras", 14, Alignment(horizontal="right", vertical="top")),
    ("Total Caracteres", 15, Alignment(horizontal="right", vertical="top")),
    ("Repetições Consecutivas", 20, Alignment(horizontal="right", vertical="top")),
    ("Taxa Repetição (%)", 18, Alignment(horizontal="right", vertical="top")),
    ("Caracteres Inválidos", 18, Alignment(horizontal="right", vertical="top")),
    ("Logprob Média", 15, Alignment(horizontal="right", vertical="top")),
    ("Taxa Compressão", 16, Alignment(horizontal="right", vertical="top")),
    ("Prob. Sem Fala Média", 18, Alignment(horizontal="right", vertical="top")),
    ("Critérios Penalizados", 40, Alignment(horizontal="left", vertical="top")),
    ("Transcrição Completa", 65, Alignment(horizontal="left", vertical="top", wrap_text=True)),
    ("Caminho Completo", 40, Alignment(horizontal="left", vertical="top")),
]


def _obter_planilha_ativa(wb: Any) -> Any:
    """Retorna a planilha ativa de forma segura, garantindo que não seja None."""
    ws: Any = wb.active
    if ws is None:
        ws = wb.create_sheet(title="Transcrições")
    return ws


def _criar_ou_carregar_workbook(caminho_arquivo: Path) -> Any:
    """Cria uma nova planilha estruturada ou abre a existente."""
    if caminho_arquivo.exists():
        try:
            wb: Any = openpyxl.load_workbook(str(caminho_arquivo))
            return wb
        except Exception:
            # Se o arquivo estiver corrompido, cria novo
            pass

    wb = openpyxl.Workbook()
    ws: Any = _obter_planilha_ativa(wb)
    ws.title = "Transcrições"

    # Adiciona cabeçalhos
    for col_idx, (nome_col, largura, _) in enumerate(COLUNAS, start=1):
        cell: Any = ws.cell(row=1, column=col_idx, value=nome_col)
        cell.fill = ESTILO_CABECALHO_FILL
        cell.font = ESTILO_CABECALHO_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDA_FINA
        try:
            letra_col = get_column_letter(col_idx)
            ws.column_dimensions[letra_col].width = largura
        except Exception:
            pass

    try:
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
    except Exception:
        pass

    return wb


def _formatar_linha(ws: Any, row_idx: int, classificacao: str) -> None:
    """Aplica bordas, alinhamentos e cores de destaque na linha inserida."""
    for col_idx, (_, _, alinhamento) in enumerate(COLUNAS, start=1):
        cell: Any = ws.cell(row=row_idx, column=col_idx)
        cell.border = BORDA_FINA
        cell.alignment = alinhamento

    # Destaque na coluna de 'Qualidade Medida' (coluna 3)
    if classificacao in CORES_CLASSIFICACAO:
        cell_qualidade: Any = ws.cell(row=row_idx, column=3)
        cell_qualidade.fill = CORES_CLASSIFICACAO[classificacao]["fill"]
        cell_qualidade.font = CORES_CLASSIFICACAO[classificacao]["font"]


def _montar_linha_dados(
    titulo_audio: str,
    caminho_audio: Path,
    texto_transcricao: str,
    dados_avaliacao: Dict[str, Any],
) -> tuple[List[Any], str]:
    """Prepara a lista com os dados ordenados para inserção na planilha."""
    agora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    classificacao = str(dados_avaliacao.get("classificacao", "Não avaliado"))
    pontuacao = dados_avaliacao.get("pontuacao", "-")
    status_fluxo = str(dados_avaliacao.get("status_fluxo", "Processado"))

    qtd_palavras = dados_avaliacao.get("quantidade_palavras", 0)
    qtd_caracteres = dados_avaliacao.get("quantidade_caracteres", len(texto_transcricao))
    qtd_repeticoes = dados_avaliacao.get("quantidade_repeticoes", 0)
    taxa_repeticao = dados_avaliacao.get("taxa_repeticao", 0.0)
    caracteres_invalidos = dados_avaliacao.get("caracteres_invalidos", 0)

    logprob = dados_avaliacao.get("logprob_media", 0.0)
    taxa_compressao = dados_avaliacao.get("taxa_compressao_media", 0.0)
    prob_sem_fala = dados_avaliacao.get("probabilidade_media_sem_fala", 0.0)

    criterios = dados_avaliacao.get("motivo")
    if not criterios and "penalizacoes" in dados_avaliacao:
        criterios = "; ".join(dados_avaliacao["penalizacoes"])
    if not criterios:
        criterios = "Nenhuma penalização"

    linha = [
        agora,
        titulo_audio,
        classificacao,
        pontuacao,
        status_fluxo,
        qtd_palavras,
        qtd_caracteres,
        qtd_repeticoes,
        round(taxa_repeticao, 2) if isinstance(taxa_repeticao, (int, float)) else taxa_repeticao,
        caracteres_invalidos,
        round(logprob, 3) if isinstance(logprob, (int, float)) else logprob,
        round(taxa_compressao, 3) if isinstance(taxa_compressao, (int, float)) else taxa_compressao,
        round(prob_sem_fala, 3) if isinstance(prob_sem_fala, (int, float)) else prob_sem_fala,
        criterios,
        texto_transcricao,
        str(caminho_audio.resolve()),
    ]
    return linha, classificacao


def exportar_transcricao_para_planilha(
    caminho_audio: Union[str, Path],
    texto_transcricao: str,
    dados_avaliacao: Dict[str, Any],
    caminho_planilha: Optional[Union[str, Path]] = None,
) -> Path:
    """
    Exporta os dados de uma transcrição e suas métricas de qualidade para a planilha.

    Parâmetros:
        caminho_audio: Caminho ou Path do arquivo de áudio transcrito.
        texto_transcricao: Texto completo gerado pelo Whisper.
        dados_avaliacao: Dicionário contendo as métricas de texto e do Whisper,
                         pontuação, classificação e penalizações.
        caminho_planilha: Caminho customizado para a planilha .xlsx (opcional).

    Retorna:
        Path do arquivo onde os dados foram salvos.
    """
    caminho_audio_path = Path(caminho_audio)
    destino = Path(caminho_planilha) if caminho_planilha else CAMINHO_PLANILHA_PADRAO
    destino.parent.mkdir(parents=True, exist_ok=True)

    linha_dados, classificacao = _montar_linha_dados(
        titulo_audio=caminho_audio_path.name,
        caminho_audio=caminho_audio_path,
        texto_transcricao=texto_transcricao,
        dados_avaliacao=dados_avaliacao,
    )

    wb: Any = _criar_ou_carregar_workbook(destino)
    ws: Any = _obter_planilha_ativa(wb)

    max_linha: int = int(ws.max_row) if ws.max_row is not None else 1
    row_idx = max_linha + 1

    ws.append(linha_dados)
    _formatar_linha(ws, row_idx, classificacao)

    try:
        if hasattr(ws, "auto_filter") and ws.auto_filter is not None:
            letra_fim = get_column_letter(len(COLUNAS))
            ws.auto_filter.ref = f"A1:{letra_fim}{row_idx}"
    except Exception:
        pass

    try:
        wb.save(str(destino))
    except PermissionError:
        caminho_alt = destino.with_name(f"{destino.stem}_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        print(f"  [aviso planilha] O arquivo {destino.name} está aberto em outro programa.")
        print(f"  [aviso planilha] Salvando dados em: {caminho_alt.name}")
        wb.save(str(caminho_alt))
        return caminho_alt

    return destino


def exportar_lote_transcricoes(
    itens: List[Dict[str, Any]],
    caminho_planilha: Optional[Union[str, Path]] = None,
) -> Path:
    """
    Exporta uma lista de itens transcritos para a planilha de uma só vez.
    Cada item deve conter: 'caminho_audio', 'texto_transcricao', 'dados_avaliacao'.
    """
    destino = Path(caminho_planilha) if caminho_planilha else CAMINHO_PLANILHA_PADRAO
    destino.parent.mkdir(parents=True, exist_ok=True)

    wb: Any = _criar_ou_carregar_workbook(destino)
    ws: Any = _obter_planilha_ativa(wb)

    for item in itens:
        caminho_audio_path = Path(item["caminho_audio"])
        texto = item.get("texto_transcricao", "")
        dados = item.get("dados_avaliacao", {})

        linha_dados, classificacao = _montar_linha_dados(
            titulo_audio=caminho_audio_path.name,
            caminho_audio=caminho_audio_path,
            texto_transcricao=texto,
            dados_avaliacao=dados,
        )

        max_linha: int = int(ws.max_row) if ws.max_row is not None else 1
        row_idx = max_linha + 1

        ws.append(linha_dados)
        _formatar_linha(ws, row_idx, classificacao)

    try:
        if hasattr(ws, "auto_filter") and ws.auto_filter is not None:
            letra_fim = get_column_letter(len(COLUNAS))
            max_final = int(ws.max_row) if ws.max_row is not None else 1
            ws.auto_filter.ref = f"A1:{letra_fim}{max_final}"
    except Exception:
        pass

    try:
        wb.save(str(destino))
    except PermissionError:
        caminho_alt = destino.with_name(f"{destino.stem}_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(str(caminho_alt))
        return caminho_alt

    return destino
